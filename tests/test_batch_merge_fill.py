"""批量 Excel 解析：图片/视频列合并单元格向下填充测试。"""

import unittest

from openpyxl import Workbook
from src.api.batch import _build_media_merged_fill_map


class BuildMediaMergedFillMapTests(unittest.TestCase):
    """直接测试 _build_media_merged_fill_map 纯函数。"""

    def _make_sheet(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Tasks"
        headers = [
            "task_type", "model", "duration", "aspect_ratio", "prompt",
            "image_1", "image_2", "image_3", "image_4", "image_5",
            "video_1", "video_2",
        ]
        for col_index, header in enumerate(headers, start=1):
            sheet.cell(row=1, column=col_index, value=header)
        # 行 2-4：image_1 列 (col 6) 纵向合并
        sheet.cell(row=2, column=6, value="shared.png")
        sheet.merge_cells(start_row=2, start_column=6, end_row=4, end_column=6)
        # 行 2-3：video_1 列 (col 11) 纵向合并
        sheet.cell(row=2, column=11, value="shared-video.mp4")
        sheet.merge_cells(start_row=2, start_column=11, end_row=3, end_column=11)
        # 横向合并（不应处理）
        sheet.merge_cells(start_row=5, start_column=6, end_row=5, end_column=8)
        return sheet

    def test_returns_merged_cell_mappings(self):
        sheet = self._make_sheet()
        media_columns = {6, 11, 7, 8, 9, 10, 12}
        result = _build_media_merged_fill_map(sheet, media_columns=media_columns)

        # image_1 列：行 3 和 4 映射到行 2
        self.assertEqual(result.get((3, 6)), 2)
        self.assertEqual(result.get((4, 6)), 2)
        # video_1 列：行 3 映射到行 2
        self.assertEqual(result.get((3, 11)), 2)
        # 合并区域的左上角不应出现在映射中
        self.assertNotIn((2, 6), result)
        self.assertNotIn((2, 11), result)
        # 横向合并 (row=5, col=6-8) 不应出现在结果中
        self.assertNotIn((5, 6), result)
        self.assertNotIn((5, 7), result)
        self.assertNotIn((5, 8), result)

    def test_only_media_columns_are_mapped(self):
        sheet = self._make_sheet()
        result = _build_media_merged_fill_map(sheet, media_columns={6})
        # video_1 列 (11) 不在 media_columns 中，不应有映射
        self.assertNotIn((3, 11), result)
        # image_1 列仍应有映射
        self.assertIn((3, 6), result)
        self.assertIn((4, 6), result)

    def test_empty_media_columns_returns_empty_map(self):
        sheet = self._make_sheet()
        result = _build_media_merged_fill_map(sheet, media_columns=set())
        self.assertEqual(result, {})

    def test_no_merged_cells_returns_empty_map(self):
        workbook = Workbook()
        sheet = workbook.active
        result = _build_media_merged_fill_map(sheet, media_columns={6})
        self.assertEqual(result, {})

    def test_header_row_merged_cells_are_ignored(self):
        sheet = self._make_sheet()
        sheet.merge_cells(start_row=1, start_column=6, end_row=1, end_column=8)
        result = _build_media_merged_fill_map(sheet, media_columns={6, 7, 8})
        for key in result:
            self.assertGreater(key[0], 1)


if __name__ == "__main__":
    unittest.main()