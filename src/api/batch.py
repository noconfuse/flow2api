"""Admin batch job APIs for Excel-driven video generation."""

import io
import json
import posixpath
import re
import secrets
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple
from urllib.parse import urlparse
from xml.etree import ElementTree as ET

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel, Field

from ..core.database import Database
from ..core.model_resolver import get_batch_video_model_catalog, resolve_batch_video_model
from ..core.models import BatchJob, BatchJobItem
from ..services.batch_executor import BatchExecutor
from ..services.generation_handler import MODEL_CONFIG
from ..services.token_manager import TokenManager
from .admin import verify_admin_token

router = APIRouter()

token_manager: Optional[TokenManager] = None
db: Optional[Database] = None
batch_executor: Optional[BatchExecutor] = None

EXPECTED_COLUMNS = [
    "task_type",
    "model",
    "duration",
    "aspect_ratio",
    "prompt",
    "image_1",
    "image_2",
    "image_3",
    "image_4",
    "image_5",
    "video_1",
    "video_2",
    "image_source",
]
SUPPORTED_TASK_TYPES = {"video", "edit", "image"}
IMAGE_SLOTS = [f"image_{index}" for index in range(1, 6)]
VIDEO_SLOTS = [f"video_{index}" for index in range(1, 3)]
DYNAMIC_IMAGE_SLOTS = ["image_source"]
SLOT_REF_RE = re.compile(
    r"@(?P<slot>image_[1-5]|video_[1-2]|image_source)\b",
    re.IGNORECASE,
)
ROW_DEPENDENCY_RE = re.compile(r"^row:(?P<row_key>[A-Za-z0-9_\-]+)$", re.IGNORECASE)
TASKS_SHEET_NAME = "Tasks"
OPTIONS_SHEET_NAME = "Options"
DEFAULT_DURATION_OPTIONS = ["4s", "6s", "8s", "10s"]
DEFAULT_ASPECT_RATIO_OPTIONS = ["16:9", "9:16"]
UPLOAD_ROOT = Path(__file__).resolve().parents[2] / "data" / "batch_uploads"

PKG_NS = {
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "doc_rel": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}

TASK_TYPE_DISPLAY_TO_INTERNAL = {
    "视频生成": "video",
    "视频编辑": "edit",
    "文生图": "image",
}

TASK_TYPE_INTERNAL_TO_DISPLAY = {
    "video": "视频生成",
    "edit": "视频编辑",
    "image": "文生图",
}


def _normalize_task_type(value: Any) -> str:
    raw = str(value or "").strip()
    if not raw:
        return ""
    return TASK_TYPE_DISPLAY_TO_INTERNAL.get(raw, "")


def _humanize_image_model_key(model_key: str, model_entry: Dict[str, Any]) -> str:
    key = str(model_key or "").strip()
    aspect = str(model_entry.get("aspect_ratio") or "").strip()
    upsample = str(model_entry.get("upsample") or "").strip()
    aspect_map = {
        "IMAGE_ASPECT_RATIO_LANDSCAPE": "横版",
        "IMAGE_ASPECT_RATIO_PORTRAIT": "竖版",
        "IMAGE_ASPECT_RATIO_SQUARE": "方形",
        "IMAGE_ASPECT_RATIO_LANDSCAPE_FOUR_THREE": "4:3",
        "IMAGE_ASPECT_RATIO_PORTRAIT_THREE_FOUR": "3:4",
    }
    upsample_map = {
        "UPSAMPLE_IMAGE_RESOLUTION_2K": "2K",
        "UPSAMPLE_IMAGE_RESOLUTION_4K": "4K",
    }

    if key.startswith("gemini-3.0-pro-image"):
        base = "Gemini 3.0 Pro 图片"
    elif key.startswith("imagen-4.0"):
        base = "Imagen 4.0 图片"
    else:
        base = f"{str(model_entry.get('model_name') or '').strip() or '图片模型'}"
    parts = [base]
    aspect_text = aspect_map.get(aspect, "")
    if aspect_text:
        parts.append(aspect_text)
    upsample_text = upsample_map.get(upsample, "")
    if upsample_text:
        parts.append(upsample_text)
    return " ".join(parts).strip() or key


def set_dependencies(tm: TokenManager, database: Database, executor: BatchExecutor) -> None:
    global token_manager, db, batch_executor
    token_manager = tm
    db = database
    batch_executor = executor


def _ensure_dependencies() -> Tuple[TokenManager, Database, BatchExecutor]:
    if token_manager is None or db is None or batch_executor is None:
        raise HTTPException(status_code=500, detail="Batch dependencies not initialized")
    return token_manager, db, batch_executor


def _require_openpyxl():
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError as exc:
        raise HTTPException(status_code=500, detail="服务缺少 Excel 依赖，请先重建容器镜像") from exc
    return Workbook, load_workbook, Alignment, Font, PatternFill, DataValidation


class ParsedBatchItem(BaseModel):
    row_index: int
    system_row_key: Optional[str] = None
    task_type: Optional[str] = None
    source_row: Dict[str, Optional[str]] = Field(default_factory=dict)
    normalized_payload: Optional[Dict[str, Any]] = None
    errors: List[Dict[str, str]] = Field(default_factory=list)
    warnings: List[Dict[str, str]] = Field(default_factory=list)


class CreateBatchJobRequest(BaseModel):
    file_name: Optional[str] = None
    items: List[ParsedBatchItem]


def _normalize_cell(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _error(code: str, message: str) -> Dict[str, str]:
    return {"code": code, "message": message}


def _build_system_row_key(row_index: int) -> str:
    return f"row_{row_index:04d}"


def _extract_slot_refs(prompt: str) -> Set[str]:
    return {match.group("slot").lower() for match in SLOT_REF_RE.finditer(prompt or "")}


def _guess_asset_kind(reference: str) -> str:
    normalized = str(reference or "").strip().lower()
    if normalized.startswith(("http://", "https://")):
        normalized = urlparse(normalized).path.lower()
    image_exts = (".png", ".jpg", ".jpeg", ".webp", ".gif")
    video_exts = (".mp4", ".mov", ".webm", ".mkv", ".avi", ".m4v")
    if normalized.endswith(image_exts):
        return "image"
    if normalized.endswith(video_exts):
        return "video"
    return "unknown"


def _normalize_archive_path(base_path: str, target: str) -> str:
    normalized_base = str(base_path or "").strip().lstrip("/")
    normalized_target = str(target or "").strip()
    if not normalized_target:
        return ""
    if normalized_target.startswith("/"):
        return normalized_target.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(normalized_base), normalized_target))


def _read_xml_from_archive(archive: zipfile.ZipFile, path: str) -> ET.Element:
    with archive.open(path) as handle:
        return ET.fromstring(handle.read())


def _read_relationships(archive: zipfile.ZipFile, owner_path: str) -> Dict[str, Dict[str, str]]:
    relationships_path = posixpath.join(
        posixpath.dirname(owner_path),
        "_rels",
        posixpath.basename(owner_path) + ".rels",
    )
    if relationships_path not in archive.namelist():
        return {}
    root = _read_xml_from_archive(archive, relationships_path)
    mapping: Dict[str, Dict[str, str]] = {}
    for relation in root.findall("rel:Relationship", PKG_NS):
        relation_id = str(relation.attrib.get("Id") or "").strip()
        if not relation_id:
            continue
        mapping[relation_id] = {
            "target": str(relation.attrib.get("Target") or "").strip(),
            "target_mode": str(relation.attrib.get("TargetMode") or "").strip(),
            "type": str(relation.attrib.get("Type") or "").strip(),
        }
    return mapping


def _extract_sheet_archive_path(archive: zipfile.ZipFile, sheet_name: str) -> str:
    workbook_path = "xl/workbook.xml"
    if workbook_path not in archive.namelist():
        raise HTTPException(status_code=400, detail="Excel 缺少 workbook.xml")
    workbook_root = _read_xml_from_archive(archive, workbook_path)
    workbook_rels = _read_relationships(archive, workbook_path)
    for sheet in workbook_root.findall("main:sheets/main:sheet", PKG_NS):
        if str(sheet.attrib.get("name") or "").strip() != sheet_name:
            continue
        relation_id = str(sheet.attrib.get(f"{{{PKG_NS['doc_rel']}}}id") or "").strip()
        relation = workbook_rels.get(relation_id)
        if not relation:
            break
        archive_path = _normalize_archive_path(workbook_path, relation.get("target") or "")
        if archive_path:
            return archive_path
    raise HTTPException(status_code=400, detail=f"Excel 缺少工作表: {sheet_name}")


def _stage_embedded_media_file(*, group_id: str, file_name: str, content: bytes) -> Dict[str, Any]:
    safe_name = Path(str(file_name or "").strip() or "embedded.bin").name
    target_dir = UPLOAD_ROOT / group_id
    target_dir.mkdir(parents=True, exist_ok=True)
    target_path = target_dir / f"{secrets.token_hex(6)}_{safe_name}"
    target_path.write_bytes(content)
    return {
        "file_name": safe_name,
        "file_path": str(target_path),
        "size_bytes": len(content),
        "asset_kind": _guess_asset_kind(safe_name),
        "reference_value": safe_name,
    }


def _extract_embedded_media_map(raw_bytes: bytes, *, sheet_name: str, staging_group_id: str) -> Dict[Tuple[int, int], Dict[str, Any]]:
    media_by_cell: Dict[Tuple[int, int], Dict[str, Any]] = {}
    try:
        archive = zipfile.ZipFile(io.BytesIO(raw_bytes))
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Excel 压缩包读取失败: {exc}") from exc

    with archive:
        sheet_path = _extract_sheet_archive_path(archive, sheet_name)
        sheet_relationships = _read_relationships(archive, sheet_path)
        drawing_targets: List[str] = []
        for relation in sheet_relationships.values():
            if relation.get("type", "").endswith("/drawing"):
                drawing_target = _normalize_archive_path(sheet_path, relation.get("target") or "")
                if drawing_target:
                    drawing_targets.append(drawing_target)

        for drawing_path in drawing_targets:
            if drawing_path not in archive.namelist():
                continue
            drawing_root = _read_xml_from_archive(archive, drawing_path)
            drawing_relationships = _read_relationships(archive, drawing_path)
            for anchor in drawing_root.findall("xdr:oneCellAnchor", PKG_NS) + drawing_root.findall("xdr:twoCellAnchor", PKG_NS):
                marker = anchor.find("xdr:from", PKG_NS)
                if marker is None:
                    continue
                row_node = marker.find("xdr:row", PKG_NS)
                col_node = marker.find("xdr:col", PKG_NS)
                if row_node is None or col_node is None:
                    continue
                try:
                    row_index = int(str(row_node.text or "0").strip()) + 1
                    col_index = int(str(col_node.text or "0").strip()) + 1
                except ValueError:
                    continue

                embed_relation_id = ""
                for element in anchor.iter():
                    for attr_name in (
                        f"{{{PKG_NS['doc_rel']}}}embed",
                        f"{{{PKG_NS['doc_rel']}}}link",
                        f"{{{PKG_NS['doc_rel']}}}id",
                    ):
                        relation_id = str(element.attrib.get(attr_name) or "").strip()
                        if relation_id and relation_id in drawing_relationships:
                            embed_relation_id = relation_id
                            break
                    if embed_relation_id:
                        break
                if not embed_relation_id:
                    continue

                relation = drawing_relationships.get(embed_relation_id) or {}
                target_mode = str(relation.get("target_mode") or "").strip().lower()
                target = str(relation.get("target") or "").strip()
                if not target:
                    continue

                if target_mode == "external":
                    external_entry = {
                        "source_url": target,
                        "file_name": urlparse(target).path.rsplit("/", 1)[-1] or target,
                        "asset_kind": _guess_asset_kind(target),
                        "reference_value": target,
                    }
                    media_by_cell.setdefault((row_index, col_index), external_entry)
                    continue

                archive_media_path = _normalize_archive_path(drawing_path, target)
                if archive_media_path not in archive.namelist():
                    continue
                file_name = posixpath.basename(archive_media_path)
                content = archive.read(archive_media_path)
                if not content:
                    continue
                staged_entry = _stage_embedded_media_file(
                    group_id=staging_group_id,
                    file_name=file_name,
                    content=content,
                )
                staged_entry["embedded_archive_path"] = archive_media_path
                media_by_cell.setdefault((row_index, col_index), staged_entry)
    return media_by_cell


def _read_media_cell(cell: Any) -> Optional[str]:
    hyperlink = getattr(cell, "hyperlink", None)
    target = str(getattr(hyperlink, "target", "") or "").strip()
    if target:
        return target
    return _normalize_cell(cell.value)


def _build_media_merged_fill_map(
    sheet: Any,
    *,
    media_columns: Set[int],
) -> Dict[Tuple[int, int], int]:
    """建立“合并区域内的格子 -> 左上角行号”的映射。

    仅对图片/视频列生效：Excel 中这些列允许纵向合并单元格，
    被合并的行 (row) 单元格实际 value 为 None，但需要继承合并区域左上角的内容。
    """
    fill_map: Dict[Tuple[int, int], int] = {}
    for cell_range in getattr(sheet, "merged_cells", {}).ranges:
        min_col, min_row, max_col, max_row = (
            cell_range.min_col,
            cell_range.min_row,
            cell_range.max_col,
            cell_range.max_row,
        )
        if min_row < 2:
            # 表头行不参与填充映射
            continue
        if max_col - min_col != 0:
            # 不处理横向合并；只把纵向合并向下填充
            continue
        col_index = min_col
        if col_index not in media_columns:
            continue
        for row_index in range(min_row + 1, max_row + 1):
            fill_map[(row_index, col_index)] = min_row
    return fill_map


def _build_media_entry(
    reference: Optional[str],
    expected_kind: str,
    *,
    embedded_entry: Optional[Dict[str, Any]] = None,
) -> Optional[Dict[str, Any]]:
    if embedded_entry:
        merged_entry = dict(embedded_entry)
        merged_entry["expected_kind"] = expected_kind
        merged_entry["reference_value"] = str(
            embedded_entry.get("reference_value")
            or embedded_entry.get("source_url")
            or embedded_entry.get("file_name")
            or ""
        ).strip()
        return merged_entry

    value = _normalize_cell(reference)
    if not value:
        return None
    if value.startswith(("http://", "https://")):
        asset_kind = _guess_asset_kind(value)
        return {
            "source_url": value,
            "file_name": urlparse(value).path.rsplit("/", 1)[-1] or value,
            "asset_kind": asset_kind,
            "reference_value": value,
        }
    return {
        "reference_value": value,
        "asset_kind": _guess_asset_kind(value),
        "expected_kind": expected_kind,
    }


def _validate_media_entry(
    slot_name: str,
    entry: Optional[Dict[str, Any]],
    expected_kind: str,
    parsed: ParsedBatchItem,
) -> Optional[Dict[str, Any]]:
    if entry is None:
        return None
    if entry.get("file_path"):
        asset_kind = str(entry.get("asset_kind") or "").strip()
        if asset_kind != expected_kind:
            parsed.errors.append(_error("file_type_mismatch", f"{slot_name} 不是有效的嵌入{ '图片' if expected_kind == 'image' else '视频' }"))
            return None
        return entry
    reference_value = str(entry.get("reference_value") or "").strip()
    if reference_value.startswith(("http://", "https://")):
        asset_kind = str(entry.get("asset_kind") or "").strip()
        if asset_kind != expected_kind:
            parsed.errors.append(_error("file_type_mismatch", f"{slot_name} 不是有效的{ '图片' if expected_kind == 'image' else '视频' } URL"))
            return None
        return entry
    parsed.errors.append(_error("unsupported_local_media", f"{slot_name} 当前只支持可访问的远程 URL"))
    return None


def _build_media_slot_refs(slot_entries: Dict[str, Dict[str, Any]]) -> Dict[str, str]:
    refs: Dict[str, str] = {}
    for slot_name, entry in slot_entries.items():
        ref_value = str(entry.get("reference_value") or entry.get("source_url") or "").strip()
        if ref_value:
            refs[slot_name] = ref_value
    return refs


def _build_reference_asset_specs(
    *,
    prompt_slots: Set[str],
    image_entries: Dict[str, Dict[str, Any]],
    video_entries: Dict[str, Dict[str, Any]],
) -> List[Dict[str, Any]]:
    slot_entries = {**image_entries, **video_entries}
    specs: List[Dict[str, Any]] = []
    for slot_name in [*IMAGE_SLOTS, *VIDEO_SLOTS]:
        entry = slot_entries.get(slot_name)
        if not entry:
            continue
        mode = "mention" if slot_name in prompt_slots else "attach"
        specs.append(
            {
                "slot": slot_name,
                "kind": "image" if slot_name.startswith("image_") else "video",
                "mode": mode,
                "reference_texts": [slot_name],
                "source_url": str(entry.get("source_url") or "").strip() or None,
                "file_path": str(entry.get("file_path") or "").strip() or None,
                "file_name": str(entry.get("file_name") or "").strip() or None,
                "asset_kind": str(entry.get("asset_kind") or "").strip() or None,
            }
        )
    return specs


def _validate_prompt_refs(
    prompt: str,
    slot_entries: Dict[str, Dict[str, Any]],
    parsed: ParsedBatchItem,
    dynamic_slot_names: Optional[set[str]] = None,
) -> Set[str]:
    prompt_slots = _extract_slot_refs(prompt)
    dynamic_set = set(dynamic_slot_names or [])
    for slot_name in sorted(prompt_slots):
        if slot_name not in slot_entries and slot_name not in dynamic_set:
            parsed.errors.append(_error("missing_prompt_slot", f"prompt 引用了未填写的素材槽位: @{slot_name}"))
    return prompt_slots


def _parse_image_source_dependency(value: Optional[str], parsed: ParsedBatchItem) -> Optional[Dict[str, Any]]:
    raw = _normalize_cell(value)
    if not raw:
        return None
    match = ROW_DEPENDENCY_RE.match(raw)
    if not match:
        parsed.warnings.append(_error("ignored_image_source_dependency", "image_source 仅支持 row:<row_key> 格式"))
        return None
    return {"row_key": str(match.group("row_key") or "").strip()}


def _get_batch_option_catalog() -> List[Dict[str, Any]]:
    video_catalog = [
        dict(item)
        for item in get_batch_video_model_catalog(MODEL_CONFIG)
        if isinstance(item, dict) and ("t2v" in (item.get("task_types") or []) or "edit" in (item.get("task_types") or []))
    ]
    image_catalog = [
        {
            "key": model_key,
            "label": _humanize_image_model_key(model_key, dict(model_entry or {})),
            "task_types": ["image"],
            "duration_options": [],
            "aspect_ratio_options": [],
            "variants": {},
        }
        for model_key, model_entry in MODEL_CONFIG.items()
        if isinstance(model_entry, dict) and str(model_entry.get("type") or "").strip().lower() == "image"
    ]
    return [*video_catalog, *sorted(image_catalog, key=lambda item: str(item.get("label") or ""))]


def _build_template_bytes() -> bytes:
    Workbook, _, Alignment, Font, PatternFill, DataValidation = _require_openpyxl()

    workbook = Workbook()
    tasks_sheet = workbook.active
    tasks_sheet.title = TASKS_SHEET_NAME
    options_sheet = workbook.create_sheet(OPTIONS_SHEET_NAME)

    catalog = _get_batch_option_catalog()
    model_options = [item["label"] for item in catalog]
    duration_options = sorted({value for item in catalog for value in item["duration_options"]} | set(DEFAULT_DURATION_OPTIONS))
    aspect_ratio_options = sorted({value for item in catalog for value in item["aspect_ratio_options"]} | set(DEFAULT_ASPECT_RATIO_OPTIONS))

    for col_index, header in enumerate(EXPECTED_COLUMNS, start=1):
        cell = tasks_sheet.cell(row=1, column=col_index, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="1E3A8A")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    first_video_entry = next(
        (item for item in catalog if "t2v" in (item.get("task_types") or [])),
        catalog[0] if catalog else {"label": "Veo 3.1 Fast", "duration_options": ["4s"]},
    )
    first_video_model = str(first_video_entry.get("label") or "")
    first_video_duration = next(
        (value for value in (first_video_entry.get("duration_options") or []) if value),
        "4s",
    )
    first_edit_entry = next(
        (item for item in catalog if "edit" in (item.get("task_types") or [])),
        catalog[0] if catalog else {"label": "Omni Flash Edit", "duration_options": ["4s"]},
    )
    first_edit_model = str(first_edit_entry.get("label") or "")
    first_edit_duration = next(
        (value for value in (first_edit_entry.get("duration_options") or []) if value),
        "4s",
    )
    first_image_model = next(
        (item["label"] for item in catalog if "image" in (item.get("task_types") or [])),
        model_options[0] if model_options else "gemini-3.0-pro-image-landscape",
    )

    example_rows = [
        [
            "视频生成",
            first_video_model,
            first_video_duration,
            "16:9",
            "清晨薄雾中的山谷，电影感镜头，镜头缓慢推进，光线柔和",
            None,
            None,
            None,
            None,
            None,
            None,
            None,
        ],
        [
            "视频生成",
            first_video_model,
            first_video_duration,
            "16:9",
            "让 @image_1 作为主体，参考 @image_2 的构图与 @image_3 的光影",
            "https://example.com/image-1.png",
            "https://example.com/image-2.png",
            "https://example.com/image-3.png",
            None,
            None,
            None,
            None,
        ],
        [
            "视频编辑",
            first_edit_model,
            first_edit_duration,
            "16:9",
            "基于 @video_1 生成更自然的镜头运动，保持主体一致",
            None,
            None,
            None,
            None,
            None,
            "https://example.com/video-1.mp4",
            None,
        ],
        [
            "文生图",
            first_image_model,
            None,
            "16:9",
            "生成一张干净背景的产品摄影风格图片，主体居中，柔光",
            None,
            None,
            None,
            None,
            None,
            None,
            None,
        ],
    ]
    for row_index, values in enumerate(example_rows, start=2):
        for col_index, value in enumerate(values, start=1):
            tasks_sheet.cell(row=row_index, column=col_index, value=value)

    option_columns = {
        "A": ["task_type", "视频生成", "视频编辑", "文生图"],
        "B": ["model", *model_options],
        "C": ["duration", *duration_options],
        "D": ["aspect_ratio", *aspect_ratio_options],
    }
    for column_name, values in option_columns.items():
        for row_index, value in enumerate(values, start=1):
            options_sheet[f"{column_name}{row_index}"] = value
    options_sheet.sheet_state = "hidden"

    validations = {
        "A": f"={OPTIONS_SHEET_NAME}!$A$2:$A$4",
        "B": f"={OPTIONS_SHEET_NAME}!$B$2:$B${1 + len(model_options)}" if model_options else None,
        "C": f"={OPTIONS_SHEET_NAME}!$C$2:$C${1 + len(duration_options)}",
        "D": f"={OPTIONS_SHEET_NAME}!$D$2:$D${1 + len(aspect_ratio_options)}",
    }
    for column_name, formula in validations.items():
        if not formula:
            continue
        validation = DataValidation(type="list", formula1=formula, allow_blank=True)
        tasks_sheet.add_data_validation(validation)
        validation.add(f"{column_name}2:{column_name}500")

    widths = {
        "A": 14, "B": 26, "C": 12, "D": 14, "E": 44,
        "F": 28, "G": 28, "H": 28, "I": 28, "J": 28, "K": 28, "L": 28, "M": 28,
    }
    for column_name, width in widths.items():
        tasks_sheet.column_dimensions[column_name].width = width
    tasks_sheet.freeze_panes = "A2"
    options_sheet["F1"] = "说明"
    options_sheet["F2"] = "task_type 用于选择任务类型；视频生成/视频编辑必须明确选择时长；image_* / video_* 支持填写远程 URL，或把媒体嵌入到对应槽位单元格；prompt 可引用 @image_1 ~ @video_2；image_source 列填写 row:<row_key> 时，prompt 中必须包含 @image_source 占位引用该列；后端会自动取同批次中对应行生成视频的最后一帧作为 image_source。"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _parse_excel_rows(raw_bytes: bytes) -> Tuple[List[str], List[ParsedBatchItem]]:
    _, load_workbook, _, _, _, _ = _require_openpyxl()
    try:
        workbook = load_workbook(io.BytesIO(raw_bytes), data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Excel 解析失败: {exc}") from exc

    if TASKS_SHEET_NAME not in workbook.sheetnames:
        raise HTTPException(status_code=400, detail=f"Excel 缺少工作表: {TASKS_SHEET_NAME}")
    sheet = workbook[TASKS_SHEET_NAME]
    staging_group_id = secrets.token_hex(8)
    embedded_media_map = _extract_embedded_media_map(
        raw_bytes,
        sheet_name=TASKS_SHEET_NAME,
        staging_group_id=staging_group_id,
    )
    header_cells = list(sheet.iter_rows(min_row=1, max_row=1))[0]
    columns = [_normalize_cell(cell.value) or "" for cell in header_cells]
    missing_columns = [column for column in EXPECTED_COLUMNS if column not in columns]
    if missing_columns:
        raise HTTPException(status_code=400, detail=f"Excel 缺少必要列: {', '.join(missing_columns)}")

    column_index_map = {column: columns.index(column) + 1 for column in EXPECTED_COLUMNS}
    media_column_indexes: Set[int] = {
        column_index_map[column]
        for column in (*IMAGE_SLOTS, *VIDEO_SLOTS)
        if column in column_index_map
    }
    media_merged_fill_map = _build_media_merged_fill_map(sheet, media_columns=media_column_indexes)
    items: List[ParsedBatchItem] = []
    logical_row_index = 0

    for excel_row in range(2, sheet.max_row + 1):
        source_row: Dict[str, Optional[str]] = {}
        embedded_slot_entries: Dict[str, Dict[str, Any]] = {}
        has_value = False
        for column in EXPECTED_COLUMNS:
            column_index = column_index_map[column]
            cell = sheet.cell(row=excel_row, column=column_index)
            is_media_column = column in (*IMAGE_SLOTS, *VIDEO_SLOTS)

            # 处理 image/video 列的纵向合并单元格：被合并的行 (row, col) 实际上 value 为 None，
            # 需要回退到合并区域左上角的格子，从而让多行复用同一张图。
            effective_row = excel_row
            if is_media_column:
                anchor_row = media_merged_fill_map.get((excel_row, column_index))
                if anchor_row:
                    effective_row = anchor_row

            embedded_entry = embedded_media_map.get((effective_row, column_index))
            if not embedded_entry and is_media_column and effective_row != excel_row:
                embedded_entry = embedded_media_map.get((excel_row, column_index))

            value_cell = sheet.cell(row=effective_row, column=column_index)
            value = _read_media_cell(value_cell) if is_media_column else _normalize_cell(value_cell.value)
            if embedded_entry and not value:
                value = str(embedded_entry.get("file_name") or "").strip() or None
            source_row[column] = value
            if embedded_entry and is_media_column:
                embedded_slot_entries[column] = embedded_entry
            if value or embedded_entry:
                has_value = True
        if not has_value:
            continue
        logical_row_index += 1
        items.append(
            _parse_excel_item(
                row_index=logical_row_index,
                source_row=source_row,
                embedded_slot_entries=embedded_slot_entries,
            )
        )

    return columns, items


def _parse_excel_item(
    *,
    row_index: int,
    source_row: Dict[str, Optional[str]],
    embedded_slot_entries: Optional[Dict[str, Dict[str, Any]]] = None,
) -> ParsedBatchItem:
    parsed = ParsedBatchItem(
        row_index=row_index,
        system_row_key=_build_system_row_key(row_index),
        task_type=source_row.get("task_type"),
        source_row=source_row,
    )
    embedded_slot_entries = embedded_slot_entries or {}

    normalized_task_type = _normalize_task_type(source_row.get("task_type"))
    model_family = str(source_row.get("model") or "").strip()
    duration = str(source_row.get("duration") or "").strip()
    aspect_ratio = str(source_row.get("aspect_ratio") or "").strip() or "16:9"
    prompt = str(source_row.get("prompt") or "").strip()
    raw_image_source = _normalize_cell(source_row.get("image_source"))

    image_entries: Dict[str, Dict[str, Any]] = {}
    video_entries: Dict[str, Dict[str, Any]] = {}

    if not normalized_task_type:
        parsed.errors.append(_error("missing_task_type", "任务类型不能为空"))
    elif normalized_task_type not in SUPPORTED_TASK_TYPES:
        parsed.errors.append(_error("unsupported_task_type", f"不支持的任务类型: {source_row.get('task_type') or ''}"))

    if not model_family:
        parsed.errors.append(_error("missing_model", "model 不能为空"))
        resolved_model_info = {"resolved_model": None}
    else:
        catalog = _get_batch_option_catalog()
        label_map = {
            str(item.get("label") or "").strip(): item
            for item in catalog
            if isinstance(item, dict) and str(item.get("label") or "").strip()
        }
        picked = label_map.get(model_family)

        if not isinstance(picked, dict):
            parsed.errors.append(_error("unsupported_model", f"不支持的模型: {model_family}"))
            resolved_model_info = {"resolved_model": None}
        else:
            picked_key = str(picked.get("key") or "").strip()
            picked_task_types = picked.get("task_types") if isinstance(picked.get("task_types"), list) else []
            if normalized_task_type == "image":
                if "image" not in picked_task_types and str((MODEL_CONFIG.get(picked_key) or {}).get("type") or "").strip().lower() != "image":
                    parsed.errors.append(_error("unsupported_model", "文生图请选择图片模型"))
                    resolved_model_info = {"resolved_model": None}
                else:
                    resolved_model_info = {
                        "resolved_model": picked_key,
                        "resolved_duration": None,
                        "resolved_aspect_ratio": aspect_ratio,
                        "task_type": "image",
                    }
            else:
                filled_image_count = sum(
                    1
                    for slot_name in IMAGE_SLOTS
                    if _normalize_cell(source_row.get(slot_name)) or slot_name in embedded_slot_entries
                )
                desired_video_type = "r2v" if filled_image_count > 0 else "t2v"
                family_key = picked_key
                if normalized_task_type == "edit":
                    desired_video_type = "edit"
                elif desired_video_type == "r2v":
                    family_key = family_key.replace("_t2v_", "_r2v_").replace("-t2v", "-r2v")
                elif desired_video_type == "t2v":
                    family_key = family_key.replace("_r2v_", "_t2v_").replace("-r2v", "-t2v")

                resolved_model_info = resolve_batch_video_model(
                    model_family=family_key,
                    task_type=desired_video_type,
                    duration=duration,
                    aspect_ratio=aspect_ratio,
                    model_config=MODEL_CONFIG,
                )
                resolved_model = str(resolved_model_info.get("resolved_model") or "").strip()
                if not resolved_model:
                    parsed.errors.append(_error("unsupported_model", str(resolved_model_info.get("error") or f"不支持的模型: {model_family}")))
                elif desired_video_type == "r2v":
                    max_images = int((MODEL_CONFIG.get(resolved_model) or {}).get("max_images") or 0)
                    filled_image_count = sum(
                        1
                        for slot_name in IMAGE_SLOTS
                        if _normalize_cell(source_row.get(slot_name)) or slot_name in embedded_slot_entries
                    )
                    if max_images > 0 and filled_image_count > max_images:
                        parsed.errors.append(_error("too_many_images", f"{resolved_model} 最多支持 {max_images} 张图片"))

    if not prompt:
        parsed.errors.append(_error("missing_prompt", "prompt 不能为空"))

    for slot_name in IMAGE_SLOTS:
        entry = _build_media_entry(
            source_row.get(slot_name),
            "image",
            embedded_entry=embedded_slot_entries.get(slot_name),
        )
        validated = _validate_media_entry(slot_name, entry, "image", parsed)
        if validated:
            image_entries[slot_name] = validated

    for slot_name in VIDEO_SLOTS:
        entry = _build_media_entry(
            source_row.get(slot_name),
            "video",
            embedded_entry=embedded_slot_entries.get(slot_name),
        )
        validated = _validate_media_entry(slot_name, entry, "video", parsed)
        if validated:
            video_entries[slot_name] = validated

    prompt_refs = _validate_prompt_refs(
        prompt,
        {**image_entries, **video_entries},
        parsed,
        dynamic_slot_names=set(DYNAMIC_IMAGE_SLOTS) if raw_image_source else None,
    )

    if normalized_task_type == "video":
        if not duration:
            parsed.errors.append(_error("missing_duration", "视频生成必须选择时长"))
        if video_entries:
            parsed.warnings.append(_error("ignored_video_slots", "视频生成会忽略视频槽位"))
    elif normalized_task_type == "edit":
        if not duration:
            parsed.errors.append(_error("missing_duration", "视频编辑必须选择时长"))
        if not video_entries:
            parsed.errors.append(_error("missing_videos", "视频编辑至少需要一个视频"))
        if image_entries:
            parsed.warnings.append(_error("ignored_image_slots", "视频编辑当前不使用图片槽位"))
        referenced_videos = [slot_name for slot_name in VIDEO_SLOTS if slot_name in prompt_refs]
        if len(referenced_videos) > 1:
            parsed.errors.append(_error("too_many_video_refs", "视频编辑当前只支持引用一个视频槽位"))
    elif normalized_task_type == "image":
        if image_entries or video_entries:
            parsed.warnings.append(_error("ignored_media_slots", "文生图会忽略素材槽位"))

    image_source_dependency = _parse_image_source_dependency(source_row.get("image_source"), parsed)

    if not parsed.errors:
        resolved_model = str(resolved_model_info.get("resolved_model") or "").strip()
        reference_asset_specs = _build_reference_asset_specs(
            prompt_slots=prompt_refs,
            image_entries=image_entries,
            video_entries=video_entries,
        )
        normalized_payload: Dict[str, Any] = {
            "model": resolved_model,
            "model_family": model_family,
            "duration": resolved_model_info.get("resolved_duration"),
            "aspect_ratio": resolved_model_info.get("resolved_aspect_ratio"),
            "prompt": prompt,
            "system_row_key": parsed.system_row_key,
            "media_slot_refs": _build_media_slot_refs({**image_entries, **video_entries}) if normalized_task_type != "image" else {},
            "reference_assets": reference_asset_specs if normalized_task_type != "image" else [],
        }
        if normalized_task_type == "video" and bool(image_entries):
            normalized_payload["input_images"] = [
                dict(image_entries[slot_name], slot=slot_name)
                for slot_name in IMAGE_SLOTS
                if slot_name in image_entries
            ]
        elif normalized_task_type == "edit":
            referenced_videos = [slot_name for slot_name in VIDEO_SLOTS if slot_name in prompt_refs and slot_name in video_entries]
            selected_slot = referenced_videos[0] if referenced_videos else next(iter(video_entries.keys()))
            normalized_payload["video_edit_input"] = dict(video_entries[selected_slot], slot=selected_slot)
        if image_source_dependency is not None:
            normalized_payload["image_source_dependency"] = image_source_dependency
        parsed.normalized_payload = normalized_payload

    parsed.task_type = TASK_TYPE_INTERNAL_TO_DISPLAY.get(normalized_task_type, "") or None
    return parsed


@router.get("/api/batch/template.xlsx")
@router.get("/api/batch/template")
async def download_batch_template(_token: str = Depends(verify_admin_token)):
    content = _build_template_bytes()
    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=batch-video-template.xlsx"},
    )


@router.post("/api/batch/excel/parse")
async def parse_batch_excel(
    file: UploadFile = File(...),
    _token: str = Depends(verify_admin_token),
):
    _ensure_dependencies()
    file_name = str(file.filename or "").strip()
    if not file_name.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="只支持上传 .xlsx 文件")

    raw_bytes = await file.read()
    if not raw_bytes:
        raise HTTPException(status_code=400, detail="上传文件为空")

    columns, items = _parse_excel_rows(raw_bytes)
    return {
        "success": True,
        "file_name": file_name,
        "columns": columns,
        "valid_count": sum(1 for item in items if not item.errors),
        "invalid_count": sum(1 for item in items if item.errors),
        "items": [item.model_dump(mode="json") for item in items],
    }


@router.post("/api/batch/jobs")
async def create_batch_job(
    request: CreateBatchJobRequest,
    _token: str = Depends(verify_admin_token),
):
    _, database, executor = _ensure_dependencies()
    if not request.items:
        raise HTTPException(status_code=400, detail="items 不能为空")

    job_id = f"batch_{secrets.token_hex(6)}"
    batch_items: List[BatchJobItem] = []
    queued_count = 0
    failed_count = 0

    for item in request.items:
        normalized_payload = item.normalized_payload if isinstance(item.normalized_payload, dict) else None
        error_message = "; ".join(
            str(entry.get("message") or "").strip()
            for entry in item.errors
            if isinstance(entry, dict) and str(entry.get("message") or "").strip()
        )
        status = "queued" if normalized_payload and not item.errors else "failed"
        if status == "queued":
            queued_count += 1
        else:
            failed_count += 1

        batch_items.append(
            BatchJobItem(
                job_id=job_id,
                row_index=item.row_index,
                row_id=item.system_row_key,
                task_type=str(item.task_type or "").strip().lower() or "unknown",
                normalized_payload=normalized_payload,
                status=status,
                error_code="validation_failed" if status == "failed" else None,
                error_message=error_message or None,
            )
        )

    initial_status = "queued" if queued_count > 0 else "failed"
    job = BatchJob(
        job_id=job_id,
        source_type="xlsx",
        file_name=request.file_name,
        raw_payload_text=json.dumps(request.model_dump(mode="json"), ensure_ascii=False),
        status=initial_status,
        total_count=len(batch_items),
        pending_count=queued_count,
        running_count=0,
        success_count=0,
        failed_count=failed_count,
        created_by="admin",
    )
    await database.create_batch_job(job, batch_items)

    if queued_count > 0:
        await executor.schedule(job_id)

    return {
        "success": True,
        "job_id": job_id,
        "status": initial_status,
        "queued_count": queued_count,
        "failed_count": failed_count,
    }


@router.get("/api/batch/jobs")
async def list_batch_jobs(
    limit: int = 50,
    _token: str = Depends(verify_admin_token),
):
    _, database, _ = _ensure_dependencies()
    jobs = await database.list_batch_jobs(limit=limit)
    return {"success": True, "items": [job.model_dump(mode="json") for job in jobs]}


@router.get("/api/batch/jobs/{job_id}")
async def get_batch_job_detail(
    job_id: str,
    _token: str = Depends(verify_admin_token),
):
    _, database, _ = _ensure_dependencies()
    job = await database.get_batch_job(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="Batch job not found")
    items = await database.list_batch_job_items(job_id)
    return {
        "success": True,
        "job": job.model_dump(mode="json"),
        "items": [item.model_dump(mode="json") for item in items],
    }


@router.delete("/api/batch/jobs/{job_id}")
async def delete_batch_job(
    job_id: str,
    _token: str = Depends(verify_admin_token),
):
    _, database, _ = _ensure_dependencies()
    job = await database.get_batch_job(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="Batch job not found")
    if str(job.status or "").strip().lower() == "running" or int(job.running_count or 0) > 0:
        raise HTTPException(status_code=409, detail="任务执行中，不能删除")
    await database.delete_batch_job(job_id)
    return {"success": True}
